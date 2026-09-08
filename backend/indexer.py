import os
import re
import json
from datetime import datetime
from bs4 import BeautifulSoup
from sqlalchemy.orm import Session
import zipfile

# Import parsers (wrap in try-except to handle imports robustly)
try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    import docx
except ImportError:
    docx = None

try:
    import mammoth
except ImportError:
    import subprocess
    import sys
    try:
        print("[Auto-setup] Installing mammoth HTML converter dependency...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", "mammoth"])
        import mammoth
    except Exception as e:
        print(f"[Auto-setup Warning] Failed to auto-install mammoth: {e}")
        mammoth = None

import markdown

from backend.database import SessionLocal, Document, HelpTopic, IndexLog, init_db

# Constants for paths
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CONFLUENCE_DIR = os.path.join(BASE_DIR, "863301644")
UPLOADS_DIR = os.path.join(BASE_DIR, "uploads")
UPLOADS_XPA = os.path.join(UPLOADS_DIR, "xpa")
UPLOADS_XPI = os.path.join(UPLOADS_DIR, "xpi")
UPLOADS_CLOUD = os.path.join(UPLOADS_DIR, "cloud_native")
UPLOADS_GENERAL = os.path.join(UPLOADS_DIR, "general")

HELP_FOLDER_NAMES = {
    "help", "help_samples", "whxdata", "mergedprojects", 
    "template", "template_scripts", "test_help_package", "test_help_package_xpa"
}

ASSET_FOLDER_NAMES = {
    "attachments", "images", "styles", "whxdata", "assets"
}

def is_help_path(file_path: str) -> bool:
    """Checks if a file path belongs to product help manuals rather than user documents."""
    norm = file_path.replace("\\", "/").lower()
    parts = set(norm.split("/"))
    if any(hf in parts for hf in HELP_FOLDER_NAMES):
        return True
    if any(k in norm for k in ["/help/", "/help_samples/", "/mergedprojects/", "/whxdata/", "/test_help_package/"]):
        return True
    base = os.path.basename(file_path).lower()
    if base in ["index.htm", "index.html"] and any(p in norm for p in ["/xpa/", "/xpi/", "/cloud_native/"]):
        return True
    return False

def is_asset_path(file_path: str) -> bool:
    """Checks if a file path belongs to an internal attachment or asset folder rather than a primary KB article."""
    norm = file_path.replace("\\", "/").lower()
    parts = set(norm.split("/"))
    if any(af in parts for af in ASSET_FOLDER_NAMES):
        return True
    return False

# Ensure all product upload directories exist
for p_dir in [UPLOADS_DIR, UPLOADS_XPA, UPLOADS_XPI, UPLOADS_CLOUD, UPLOADS_GENERAL]:
    os.makedirs(p_dir, exist_ok=True)

# An upload folder's NAME decides which product space owns it, by token rather
# than by an exact list. The naming convention is "<anything>-<product>" or
# "<product>-<platform>", so xpa-windows, xpa-linux and isolated-fix-xpa all
# resolve to Magic xpa; likewise for xpi. Cloud Service Ops (cso) is Cloud Native.
#
# Matching on a token instead of a fixed list is deliberate: the previous version
# only knew the bare names (xpa, xpi, cloud_native, general) — all of which are
# empty directories — so the real folders (cso/, xpi-windows/, xpa-windows/…)
# matched nothing, fell through to keyword guessing, and a CSO runbook titled
# "SOP for XPA Projects Alerts" landed in the Magic xpa space. New platform or
# isolated-fix variants are now picked up without another code change.
#
# A folder naming BOTH products (e.g. today's flat "isolated-fix" holds
# "Send Data from XPA to XPI") is left unowned on purpose, so those documents
# keep being classified per-document rather than dumped wholesale into one space.
PRODUCT_DIR_TOKENS = [
    (("cso", "cloud_native", "cloud"), "cloud_native"),
    (("xpa",), "xpa"),
    (("xpi",), "xpi"),
]


def _folder_product(folder: str):
    """Resolves one folder name to a product space, or None if it names none/both."""
    parts = [p for p in re.split(r"[^a-z0-9]+", folder.lower()) if p]
    hits = {prod for tokens, prod in PRODUCT_DIR_TOKENS if any(t in parts for t in tokens)}
    if len(hits) == 1:
        return hits.pop()
    return None                      # names no product, or is ambiguous


def product_from_path(file_path: str):
    """Returns the product space a path's folder dictates, or None if unowned.

    Only folders beneath uploads/ are considered, so a machine-specific parent
    directory that happens to contain "xpi" can never decide a product.
    """
    path_lower = (file_path or "").lower().replace("\\", "/")
    if "/uploads/" not in path_lower:
        return None
    segments = path_lower.split("/uploads/", 1)[1].split("/")[:-1]   # drop filename
    for seg in segments:
        if seg in ("attachments", "styles", "assets", "images"):
            continue                 # export scaffolding, never a product space
        if seg == "general":
            return "general"
        owned = _folder_product(seg)
        if owned:
            return owned
    return None


def detect_product(file_path: str, title: str = "", content: str = "", breadcrumbs: list = None) -> str:
    """
    Identifies which Magic product space an article belongs to:
    'xpa' (Application Platform), 'xpi' (Integration Platform), 'cloud_native', or 'general'.
    Physical directory structure is strictly authoritative.
    """
    # 1. Folder path is strictly authoritative when placed inside product directories
    owned = product_from_path(file_path)
    if owned:
        return owned

    crumbs_text = " ".join(breadcrumbs) if breadcrumbs else ""
    # 2. Confluence breadcrumbs check (e.g. Cloud Service Ops in root Confluence export)
    if any(k in crumbs_text.lower() for k in ["cloud service ops", "cloud native", "cso"]):
        return "cloud_native"

    combined = (title + " " + crumbs_text + " " + (content[:3500] if content else "")).lower()
    
    # 3. Content heuristics for ambiguous root dumps (e.g. 863301644)
    # NB: "imm" is deliberately NOT a Cloud Native keyword. IMM is a Magic xpi
    # 4.14 component, so its install/troubleshooting articles ("Unable to Start
    # IMM Agent", "4.14.1_IMM_VHDX installation guide") belong to xpi. An IMM
    # article that is genuinely about container orchestration still lands in
    # Cloud Native via the kubernetes/helm/docker keywords below.
    if any(k in combined for k in ["cloud native", "kubernetes", "k8s", "docker", "helm", "microservice", "modernization factory", "cloud migration"]):
        return "cloud_native"
        
    if any(k in combined for k in ["xpi", "gigaspace", "gsc", "gsm", "gsa", "datamapper", "data mapper", "connector", "sap b1", "salesforce connector", "sugarcrm", "dynamics", "tcp-listener", "http trigger", "rest client"]):
        return "xpi"
        
    if any(k in combined for k in ["magic xpa", "xpa studio", "ria application", "unipaas", "magic.ini", "mgreq", "mgrb", "xpa "]):
        if "xpi" not in combined:
            return "xpa"

    if "863301644" in file_path:
        return "xpi"
        
    return "general"

def detect_version(title: str = "", content: str = "") -> str:
    """Extracts software version from title or first paragraphs."""
    text = title + " " + content[:600]
    match = re.search(r'\b(4\.\d+(\.\d+)?|3\.\d+|12\.\d+|10\.\d+|v\d+(\.\d+)?)\b', text, re.IGNORECASE)
    if match:
        return match.group(0)
    return "Universal"

def detect_doc_type(title: str = "", content: str = "", breadcrumbs: list = None) -> str:
    """Categorizes document into function, syntax, error_code, troubleshooting, how_to, connector, architecture, release_note."""
    crumbs_str = " ".join(breadcrumbs or []).lower()
    text = (title + " " + crumbs_str + " " + content[:500]).lower()
    if any(w in crumbs_str for w in ["function directory", "expression editor", "functions", "function list"]):
        return "function"
    if "syntax:" in text[:300] and any(w in text[:300] for w in ["parameters:", "returns:", "example:"]):
        return "function"
    if any(w in crumbs_str for w in ["error code", "error codes", "error messages"]) or "error code" in title.lower():
        return "error_code"
    if any(w in text for w in ["error", "issue", "problem", "fail", "leak", "warning", "fix", "troubleshoot", "crash", "bug"]):
        return "troubleshooting"
    if any(w in text for w in ["how to", "how-to", "guide", "setup", "configure", "installation", "steps"]):
        return "how_to"
    if any(w in text for w in ["connector", "adapter", "interface", "api", "rest", "soap", "odata", "salesforce", "sap"]):
        return "connector"
    if any(w in text for w in ["architecture", "overview", "lifecycle", "design", "specification"]):
        return "architecture"
    if any(w in text for w in ["release", "what's new", "changelog", "version"]):
        return "release_note"
    return "general"

def parse_html_file(file_path):
    """
    Parses Confluence HTML and RoboHelp/WebHelp topic files.
    Extracts: title, breadcrumbs, syntax signature, author, date, and clean body text.
    """
    try:
        norm_path = file_path.lower().replace("\\", "/")
        # Filter out JavaScript/CSS template and navigation internal assets
        if any(part in norm_path for part in ["/whxdata/", "/template/", "/scripts/", "/template_scripts/"]):
            return None
        # Filter out root frameset index files that contain no body content
        base_name = os.path.basename(file_path).lower()
        if base_name in ["index.htm", "index.html"]:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content_sample = f.read(4000).lower()
                if "gtopicframename" in content_sample or "modernlayoutcontroller" in content_sample or "frameset" in content_sample:
                    return None

        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            soup = BeautifulSoup(f.read(), "html.parser")

        # 1. Breadcrumbs
        breadcrumbs = []
        meta_bc = soup.find("meta", attrs={"name": re.compile(r'topic-breadcrumbs|breadcrumbs', re.I)})
        if meta_bc and meta_bc.get("content"):
            breadcrumbs = [b.strip() for b in meta_bc["content"].split(">") if b.strip()]
        else:
            bc_section = soup.find(id="breadcrumbs") or soup.find(class_="breadcrumbs")
            if bc_section:
                breadcrumbs = [li.get_text().strip() for li in bc_section.find_all(["li", "a"]) if li.get_text().strip()]

        # 2. Title
        title = ""
        # Check H1 first (RoboHelp topics place clean function/topic names in H1)
        h1_tag = soup.find("h1")
        if h1_tag:
            title = h1_tag.get_text().strip()
        
        if not title:
            title_tag = soup.find(id="title-text")
            if title_tag:
                title = title_tag.get_text().strip()
            else:
                title_tag = soup.find("title")
                if title_tag:
                    title = title_tag.get_text().strip()
        
        if not title or title.lower() in ["untitled document", "magic xpa help", "magic xpi help", "help"]:
            title = os.path.splitext(os.path.basename(file_path))[0].replace("_", " ")

        # Clean confluence/help prefix
        if "Magic Global Support : " in title:
            title = title.replace("Magic Global Support : ", "")
        title = re.sub(r'^(Magic\s+(xpa|xpi)\s+[\d\.]*\s*Help\s*-\s*)', '', title, flags=re.I).strip()

        # 3. Syntax Extraction (for function reference topics)
        syntax_str = ""
        syntax_cell = soup.find(lambda e: e.name in ["td", "th", "p", "dt", "b", "strong"] and "syntax:" in e.get_text().lower())
        if syntax_cell:
            if syntax_cell.name in ["td", "th"]:
                sibling = syntax_cell.find_next_sibling(["td", "th"])
                if sibling:
                    syntax_str = sibling.get_text().strip()
            if not syntax_str:
                cell_text = syntax_cell.get_text().strip()
                match = re.search(r'syntax:\s*([^\n\r]+)', cell_text, re.I)
                if match:
                    syntax_str = match.group(1).strip()
            if not syntax_str:
                next_p = syntax_cell.find_next(["p", "div", "pre", "code"])
                if next_p:
                    syntax_str = next_p.get_text().strip()

        # 4. Author and Date
        author = "Magic Documentation"
        doc_date = None
        meta_div = soup.find(class_="page-metadata")
        if meta_div:
            author_span = meta_div.find(class_="author")
            if author_span:
                author = author_span.get_text().strip()
            meta_text = meta_div.get_text()
            date_match = re.search(r"on\s+([0-9a-zA-Z,\s]+)", meta_text)
            if date_match:
                date_str = date_match.group(1).strip()
                try:
                    date_cleaned = date_str.replace(" ", "").replace(",", "")
                    doc_date = datetime.strptime(date_cleaned, "%d%b%Y")
                except Exception:
                    pass

        # 5. Content
        content_div = soup.find(id="main-content") or soup.find(class_="wiki-content") or soup.find("body")
        if content_div:
            for element in content_div(["script", "style", "nav", "header", "footer"]):
                element.decompose()
            content = content_div.get_text(separator=" ").strip()
        else:
            content = ""

        content = re.sub(r"\s+", " ", content)
        if len(content) < 15 and not syntax_str:
            return None

        tags_list = []
        if syntax_str:
            tags_list.append(f"syntax:{syntax_str}")
        crumbs_str = " ".join(breadcrumbs).lower()
        if any(w in crumbs_str for w in ["function directory", "expression editor", "functions"]) or syntax_str:
            tags_list.extend(["function", "syntax", "reference", title.lower()])
        elif any(w in crumbs_str for w in ["error code", "error codes"]):
            tags_list.extend(["error_code", title.lower()])

        return {
            "title": title,
            "breadcrumbs": json.dumps(breadcrumbs),
            "author": author,
            "created_at": doc_date or datetime.fromtimestamp(os.path.getmtime(file_path)),
            "content": content,
            "tags": ";".join(tags_list) if tags_list else None,
            "syntax": syntax_str
        }
    except Exception as e:
        print(f"Error parsing HTML {file_path}: {e}")
        return None

def extract_docx_images(file_path, doc_dir_name):
    """
    Extracts all images from a docx file and saves them to attachments/{doc_dir_name}/
    """
    image_paths = []
    try:
        dest_dir = os.path.join(CONFLUENCE_DIR, "attachments", doc_dir_name)
        os.makedirs(dest_dir, exist_ok=True)
        
        with zipfile.ZipFile(file_path, 'r') as zp:
            for item in zp.infolist():
                if item.filename.startswith('word/media/'):
                    ext = item.filename.split('.')[-1].lower()
                    if ext in ['png', 'jpg', 'jpeg', 'gif']:
                        img_name = os.path.basename(item.filename)
                        dest_path = os.path.join(dest_dir, img_name)
                        with open(dest_path, 'wb') as f:
                            f.write(zp.read(item.filename))
                        image_paths.append(f"attachments/{doc_dir_name}/{img_name}")
    except Exception as e:
        print(f"Error extracting docx images: {e}")
    return image_paths

def extract_pdf_images(file_path, doc_dir_name):
    """
    Extracts all images from a PDF file using pypdf and saves them to attachments/{doc_dir_name}/
    """
    image_paths = []
    if not PdfReader:
        return image_paths
    try:
        reader = PdfReader(file_path)
        dest_dir = os.path.join(CONFLUENCE_DIR, "attachments", doc_dir_name)
        os.makedirs(dest_dir, exist_ok=True)
        
        img_idx = 1
        for page_idx, page in enumerate(reader.pages):
            if "/Resources" in page and "/XObject" in page["/Resources"]:
                xObject = page["/Resources"]["/XObject"].get_object()
                for obj_name in xObject:
                    if xObject[obj_name]["/Subtype"] == "/Image":
                        data = xObject[obj_name].get_data()
                        if "/Filter" in xObject[obj_name]:
                            ext = "png"
                            filt = xObject[obj_name]["/Filter"]
                            if isinstance(filt, list):
                                filt = filt[0] if len(filt) > 0 else ""
                            
                            if "/DCTDecode" in str(filt):
                                ext = "jpg"
                            elif "/JPXDecode" in str(filt):
                                ext = "jp2"
                            
                            img_name = f"page_{page_idx+1}_img_{img_idx}.{ext}"
                            dest_path = os.path.join(dest_dir, img_name)
                            with open(dest_path, "wb") as f:
                                f.write(data)
                            image_paths.append(f"attachments/{doc_dir_name}/{img_name}")
                            img_idx += 1
    except Exception as e:
        print(f"Error extracting PDF images: {e}")
    return image_paths

def parse_pdf_file(file_path):
    """
    Parses a PDF file using pypdf.
    """
    if not PdfReader:
        print(f"pypdf not installed. Skipping extraction for {file_path}")
        return None
        
    try:
        reader = PdfReader(file_path)
        content = []
        for page in reader.pages:
            # Avoid extraction_mode="layout" to prevent character splitting spaces (e.g. confi guration)
            text = page.extract_text()
            if text:
                content.append(text)
                
        # Join pages with clean spacing
        full_content = "\n\n".join(content)
        full_content = full_content.replace("\r", "")
        
        # Clean double spaces but keep newlines
        lines = []
        for line in full_content.split("\n"):
            cleaned_line = re.sub(r"[ \t]+", " ", line).strip()
            lines.append(cleaned_line)
        full_content = "\n".join(lines)
        
        title = os.path.basename(file_path)
        if reader.metadata and reader.metadata.title:
            title = reader.metadata.title
            
        # Extract and save images dynamically
        doc_filename = os.path.basename(file_path)
        doc_dir_name = re.sub(r'\W+', '_', doc_filename)
        extract_pdf_images(file_path, doc_dir_name)
            
        return {
            "title": title,
            "breadcrumbs": json.dumps(["Uploads", "PDFs"]),
            "author": reader.metadata.creator if reader.metadata and reader.metadata.creator else "Editor",
            "created_at": datetime.fromtimestamp(os.path.getmtime(file_path)),
            "content": full_content.strip()
        }
    except Exception as e:
        print(f"Error parsing PDF {file_path}: {e}")
        return None

def parse_docx_file(file_path):
    """
    Parses a Word DOCX file.
    If mammoth is available, converts it to clean HTML (preserving tables, images) and saves it on disk.
    """
    if not docx:
        print(f"python-docx not installed. Skipping extraction for {file_path}")
        return None
        
    try:
        # 1. Extract raw text for DB content search (using python-docx)
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs]
        full_content = "\n\n".join(paragraphs)
        
        title = os.path.basename(file_path)
        if len(doc.paragraphs) > 0 and doc.paragraphs[0].style.name.startswith("Heading"):
            title = doc.paragraphs[0].text
            
        # 2. Try converting to high-fidelity HTML with Mammoth
        if mammoth:
            try:
                with open(file_path, "rb") as docx_file:
                    result = mammoth.convert_to_html(docx_file)
                    html_val = result.value
                    
                    # Style tables beautifully inline
                    html_val = html_val.replace("<table>", '<table style="width:100%; border-collapse:collapse; margin-bottom:1.5rem; background:rgba(255,255,255,0.01); border:1px solid rgba(255,255,255,0.08);">')
                    html_val = html_val.replace("<th>", '<th style="border:1px solid rgba(255,255,255,0.08); padding:10px 12px; background:rgba(255,255,255,0.04); text-align:left; font-weight:600; color:#fff; font-size:0.88rem;">')
                    html_val = html_val.replace("<td>", '<td style="border:1px solid rgba(255,255,255,0.08); padding:10px 12px; color:#d1d5db; font-size:0.88rem; line-height:1.6;">')
                    
                    # Wrap lists in clean style
                    html_val = html_val.replace("<ul>", '<ul style="margin-left:1.8rem; margin-bottom:1.2rem; list-style-type:disc; color:#d1d5db;">')
                    html_val = html_val.replace("<ol>", '<ol style="margin-left:1.8rem; margin-bottom:1.2rem; list-style-type:decimal; color:#d1d5db;">')
                    html_val = html_val.replace("<li>", '<li style="margin-bottom:0.5rem; line-height:1.7;">')
                    
                    # Style images to be responsive
                    html_val = html_val.replace("<img ", '<img style="max-width:100%; height:auto; border-radius:8px; margin:16px 0; box-shadow:0 4px 15px rgba(0,0,0,0.3); border:1px solid rgba(255,255,255,0.08);" ')
                    
                    # Save HTML file next to original docx file
                    html_path = file_path + ".html"
                    with open(html_path, "w", encoding="utf-8") as hf:
                        hf.write(html_val)
                    print(f"Mammoth HTML generated and saved: {html_path}")
            except Exception as e:
                print(f"Mammoth conversion failed for {file_path}: {e}")
                
        # Also extract individual images to attachments directory for reference
        doc_filename = os.path.basename(file_path)
        doc_dir_name = re.sub(r'\W+', '_', doc_filename)
        extract_docx_images(file_path, doc_dir_name)
            
        return {
            "title": title,
            "breadcrumbs": json.dumps(["Uploads", "DOCX"]),
            "author": "Editor",
            "created_at": datetime.fromtimestamp(os.path.getmtime(file_path)),
            "content": full_content.strip()
        }
    except Exception as e:
        print(f"Error parsing DOCX {file_path}: {e}")
        return None

def parse_text_file(file_path, file_type):
    """
    Parses plain text or Markdown files.
    """
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text_content = f.read()
            
        title = os.path.basename(file_path)
        content_for_db = text_content
        
        # If Markdown, extract first heading as title
        if file_type == "md":
            html_content = markdown.markdown(text_content)
            soup = BeautifulSoup(html_content, "html.parser")
            h1 = soup.find("h1")
            if h1:
                title = h1.get_text().strip()
            content_for_db = soup.get_text()
            
        # Clean double spaces but keep paragraph layouts
        lines = []
        for line in content_for_db.split("\n"):
            cleaned_line = re.sub(r"[ \t]+", " ", line).strip()
            lines.append(cleaned_line)
        full_content = "\n".join(lines)
        
        return {
            "title": title,
            "breadcrumbs": json.dumps(["Uploads", file_type.upper()]),
            "author": "Editor",
            "created_at": datetime.fromtimestamp(os.path.getmtime(file_path)),
            "content": full_content.strip()
        }
    except Exception as e:
        print(f"Error parsing text file {file_path}: {e}")
        return None

def clean_and_migrate_help_docs(db: Session) -> int:
    """
    Cleans up any auto-indexed Help manual HTML topics from the documents table
    so the Space Documents tree and homepage counts strictly reflect real user documents.
    """
    cleaned = 0
    try:
        all_docs = db.query(Document).all()
        for doc in all_docs:
            if is_help_path(doc.file_path):
                db.delete(doc)
                cleaned += 1
        if cleaned > 0:
            db.commit()
            print(f"[Help Cleanup] Purged {cleaned} Help manual files from regular documents table.")
    except Exception as e:
        db.rollback()
        print(f"[Help Cleanup Warning] {e}")
    return cleaned

def scan_and_index_help_topics(db: Session) -> int:
    """
    Dedicated high-speed indexer for official MSE product Help manuals (RoboHelp/WebHelp topics).
    Extracts function signatures, methods, syntax, parameter specs, and how-tos into HelpTopic.
    """
    indexed_help_count = 0
    help_roots = [
        ("xpa", os.path.join(UPLOADS_XPA, "Help")),
        ("xpi", os.path.join(UPLOADS_XPI, "Help")),
    ]
    if os.path.exists(os.path.join(UPLOADS_CLOUD, "Help")):
        help_roots.append(("cloud_native", os.path.join(UPLOADS_CLOUD, "Help")))
    if os.path.exists(os.path.join(UPLOADS_XPA, "help_samples")):
        help_roots.append(("xpa", os.path.join(UPLOADS_XPA, "help_samples")))

    existing_map = {t.file_path: t for t in db.query(HelpTopic).all()}

    for default_product, help_dir in help_roots:
        if not os.path.exists(help_dir):
            continue
            
        for root, dirs, files in os.walk(help_dir):
            norm_root = root.replace("\\", "/").lower()
            if any(p in norm_root for p in ["/whxdata", "/template", "/scripts", "/template_scripts"]):
                continue
                
            for file in files:
                ext = file.split(".")[-1].lower()
                if ext not in ["html", "htm"]:
                    continue
                file_path = os.path.join(root, file)
                normalized_path = os.path.abspath(file_path)

                mtime = datetime.fromtimestamp(os.path.getmtime(file_path))
                existing = existing_map.get(normalized_path)
                if existing and existing.created_at >= mtime:
                    continue

                parsed = parse_html_file(file_path)
                if not parsed or not parsed.get("content"):
                    continue

                title = parsed["title"]
                syntax_str = parsed.get("syntax") or ""
                crumbs_json = parsed.get("breadcrumbs") or "[]"
                
                doc_type = "reference"
                if syntax_str or "syntax" in (parsed.get("tags") or ""):
                    doc_type = "function"
                elif "how to" in title.lower() or "getting started" in title.lower():
                    doc_type = "how_to"

                if existing:
                    existing.title = title
                    existing.content = parsed["content"]
                    existing.syntax = syntax_str
                    existing.breadcrumbs = crumbs_json
                    existing.doc_type = doc_type
                    existing.created_at = mtime
                else:
                    new_topic = HelpTopic(
                        title=title,
                        product=default_product,
                        file_path=normalized_path,
                        file_type=ext,
                        content=parsed["content"],
                        syntax=syntax_str,
                        breadcrumbs=crumbs_json,
                        doc_type=doc_type,
                        created_at=mtime
                    )
                    db.add(new_topic)
                    existing_map[normalized_path] = new_topic
                indexed_help_count += 1
                if indexed_help_count % 200 == 0:
                    db.commit()

    db.commit()
    print(f"[Help Indexer] Total {len(existing_map)} official help topics active. Updated: {indexed_help_count}")
    return len(existing_map)

def scan_and_index(db: Session):
    """
    Scans directory, parses files, inserts/updates DB entries, and deletes stale ones.
    """
    log_msg = []
    indexed_count = 0
    start_time = datetime.utcnow()
    
    # Initialize DB tables
    init_db()
    
    # 1. Clean and migrate any help manual topics previously stored in documents table
    clean_and_migrate_help_docs(db)

    # 1b. Clean up any internal attachment artifacts previously indexed as standalone documents
    stale_attachments = db.query(Document).filter(
        (Document.file_path.like("%/attachments/%")) | 
        (Document.file_path.like("%\\attachments\\%")) |
        (Document.file_path.like("%/images/%")) | 
        (Document.file_path.like("%\\images\\%"))
    ).all()
    if stale_attachments:
        for sa in stale_attachments:
            db.delete(sa)
        db.commit()
        log_msg.append(f"Cleaned up {len(stale_attachments)} internal attachment artifacts from KB index.")

    # 1c. Ensure all existing documents match their authoritative folder space
    aligned_count = 0
    for doc in db.query(Document).all():
        crumbs = doc.breadcrumbs or ""
        # Same single source of truth as detect_product(); the folder wins, and
        # Confluence breadcrumbs only decide for files outside a product folder.
        target_prod = product_from_path(doc.file_path)
        if not target_prod and "cloud service ops" in crumbs.lower():
            target_prod = "cloud_native"
        if target_prod and doc.product != target_prod:
            doc.product = target_prod
            aligned_count += 1
    if aligned_count:
        db.commit()
        log_msg.append(f"Re-aligned {aligned_count} documents to their authoritative folder space.")

    # 2. Scan for real Knowledge Base files
    all_files = []
    
    # Confluence folder (discover all supported file types: html, htm, docx, pdf, txt, md)
    if os.path.exists(CONFLUENCE_DIR):
        for root, dirs, files in os.walk(CONFLUENCE_DIR):
            dirs[:] = [d for d in dirs if d.lower() not in HELP_FOLDER_NAMES and d.lower() not in ASSET_FOLDER_NAMES]
            if is_help_path(root) or is_asset_path(root):
                continue
            for file in files:
                if file in ["index.html", "index.htm"]:
                    continue
                file_full = os.path.join(root, file)
                if is_help_path(file_full) or is_asset_path(file_full):
                    continue
                ext = file.split(".")[-1].lower()
                if ext in ["html", "htm", "pdf", "docx", "txt", "md"]:
                    all_files.append((file_full, ext))
                    
    # Uploads folder (and all product subdirectories)
    if os.path.exists(UPLOADS_DIR):
        for root, dirs, files in os.walk(UPLOADS_DIR):
            dirs[:] = [d for d in dirs if d.lower() not in HELP_FOLDER_NAMES and d.lower() not in ASSET_FOLDER_NAMES]
            if is_help_path(root) or is_asset_path(root):
                continue
            for file in files:
                if file in ["index.html", "index.htm"]:
                    continue
                file_full = os.path.join(root, file)
                if is_help_path(file_full) or is_asset_path(file_full):
                    continue
                ext = file.split(".")[-1].lower()
                if ext in ["html", "htm", "pdf", "docx", "txt", "md"]:
                    all_files.append((file_full, ext))

    log_msg.append(f"Discovered {len(all_files)} real KB documents on disk.")
    
    # Track paths processed to identify deleted ones
    processed_paths = set()
    
    # 2. Index each file
    for file_path, file_type in all_files:
        normalized_path = os.path.abspath(file_path)
        processed_paths.add(normalized_path)
        
        # Check if already in DB and up to date (compare modification time)
        existing_doc = db.query(Document).filter(Document.file_path == normalized_path).first()
        mtime = datetime.fromtimestamp(os.path.getmtime(file_path))
        
        # Force re-parse if the document contains flattened text (no newlines) from the previous indexer version,
        # or if it is a DOCX file but the Mammoth HTML cache does not exist on disk yet
        needs_reparse = False
        if existing_doc:
            if existing_doc.file_type in ["pdf", "docx", "txt", "md"] and "\n" not in existing_doc.content:
                needs_reparse = True
            elif existing_doc.file_type == "docx" and not os.path.exists(normalized_path + ".html"):
                needs_reparse = True
            
        if existing_doc and not needs_reparse and existing_doc.created_at >= mtime:
            # Document exists and is up to date
            indexed_count += 1
            continue
            
        # Parse based on file type
        parsed_data = None
        if file_type in ["html", "htm"]:
            parsed_data = parse_html_file(file_path)
        elif file_type == "pdf":
            parsed_data = parse_pdf_file(file_path)
        elif file_type == "docx":
            parsed_data = parse_docx_file(file_path)
        elif file_type in ["txt", "md"]:
            parsed_data = parse_text_file(file_path, file_type)
            
        if parsed_data:
            breadcrumbs_list = json.loads(parsed_data["breadcrumbs"]) if parsed_data["breadcrumbs"] else []
            product_tag = detect_product(file_path, parsed_data["title"], parsed_data["content"], breadcrumbs_list)
            version_tag = detect_version(parsed_data["title"], parsed_data["content"])
            doc_type_tag = detect_doc_type(parsed_data["title"], parsed_data["content"], breadcrumbs_list)
            
            if existing_doc:
                existing_doc.title = parsed_data["title"]
                existing_doc.breadcrumbs = parsed_data["breadcrumbs"]
                existing_doc.content = parsed_data["content"]
                if parsed_data.get("tags"):
                    existing_doc.tags = parsed_data["tags"]
                if parsed_data.get("author") and parsed_data["author"] not in ["Editor", "System", "Magic Documentation"] or not existing_doc.author:
                    existing_doc.author = parsed_data["author"]
                existing_doc.created_at = parsed_data["created_at"]
                if not existing_doc.product or existing_doc.product == "xpi":
                    existing_doc.product = product_tag
                if not existing_doc.version or existing_doc.version == "Universal":
                    existing_doc.version = version_tag
                if not existing_doc.doc_type:
                    existing_doc.doc_type = doc_type_tag
            else:
                # Create new
                new_doc = Document(
                    title=parsed_data["title"],
                    file_path=normalized_path,
                    file_type=file_type,
                    content=parsed_data["content"],
                    author=parsed_data["author"],
                    breadcrumbs=parsed_data["breadcrumbs"],
                    product=product_tag,
                    version=version_tag,
                    doc_type=doc_type_tag,
                    tags=parsed_data.get("tags"),
                    status="published",
                    created_at=parsed_data["created_at"]
                )
                db.add(new_doc)
            
            indexed_count += 1
            db.commit()
            
    # 3. Clean up deleted files from DB
    db_docs = db.query(Document).all()
    deleted_count = 0
    for doc in db_docs:
        if doc.file_path not in processed_paths:
            db.delete(doc)
            deleted_count += 1
            
    if deleted_count > 0:
        db.commit()
        log_msg.append(f"Removed {deleted_count} stale documents from search index.")
        
    # Also index official product Help topics into dedicated HelpTopic storage
    try:
        scan_and_index_help_topics(db)
    except Exception as e:
        print(f"[Help Indexing Error] {e}")

    end_time = datetime.utcnow()
    duration = (end_time - start_time).total_seconds()
    log_msg.append(f"Successfully processed search index. Total documents indexed: {indexed_count}.")
    log_msg.append(f"Indexing completed in {duration:.2f} seconds.")
    
    # 4. Log indexing run
    final_message = "\n".join(log_msg)
    idx_log = IndexLog(
        timestamp=start_time,
        status="Success",
        message=final_message,
        indexed_count=indexed_count
    )
    db.add(idx_log)
    db.commit()
    print(final_message)
    return indexed_count, final_message

def index_single_file(file_path: str, file_type: str, db: Session, explicit_product: str = None, explicit_version: str = None, explicit_doc_type: str = None):
    """
    Instantly parses and indexes a single uploaded file into the database in under 50ms.
    Eliminates full workspace re-scanning delays during uploads.
    """
    normalized_path = os.path.abspath(file_path)
    
    parsed_data = None
    file_type = file_type.lower()
    if file_type in ["html", "htm"]:
        parsed_data = parse_html_file(file_path)
    elif file_type == "pdf":
        parsed_data = parse_pdf_file(file_path)
    elif file_type == "docx":
        parsed_data = parse_docx_file(file_path)
    elif file_type in ["txt", "md"]:
        parsed_data = parse_text_file(file_path, file_type)
        
    if not parsed_data:
        return None

    breadcrumbs_list = json.loads(parsed_data["breadcrumbs"]) if parsed_data["breadcrumbs"] else []
    product_tag = explicit_product or detect_product(file_path, parsed_data["title"], parsed_data["content"], breadcrumbs_list)
    version_tag = explicit_version or detect_version(parsed_data["title"], parsed_data["content"])
    doc_type_tag = explicit_doc_type or detect_doc_type(parsed_data["title"], parsed_data["content"], breadcrumbs_list)

    existing_doc = db.query(Document).filter(Document.file_path == normalized_path).first()
    if existing_doc:
        existing_doc.title = parsed_data["title"]
        existing_doc.breadcrumbs = parsed_data["breadcrumbs"]
        existing_doc.author = parsed_data["author"]
        existing_doc.content = parsed_data["content"]
        existing_doc.product = product_tag
        existing_doc.version = version_tag
        existing_doc.doc_type = doc_type_tag
        if parsed_data.get("tags"):
            existing_doc.tags = parsed_data["tags"]
        existing_doc.created_at = parsed_data["created_at"]
        doc = existing_doc
    else:
        doc = Document(
            title=parsed_data["title"],
            file_path=normalized_path,
            file_type=file_type,
            content=parsed_data["content"],
            author=parsed_data["author"],
            breadcrumbs=parsed_data["breadcrumbs"],
            product=product_tag,
            version=version_tag,
            doc_type=doc_type_tag,
            tags=parsed_data.get("tags"),
            status="published",
            created_at=parsed_data["created_at"]
        )
        db.add(doc)
        
    db.commit()
    db.refresh(doc)
    return doc

if __name__ == "__main__":
    db = SessionLocal()
    try:
        scan_and_index(db)
    finally:
        db.close()
